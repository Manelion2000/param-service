from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = ROOT / "docs" / "user-manual"
SCREENSHOTS = OUT_DIR / "screenshots"
DOCX_PATH = ROOT / "docs" / "Manuel_utilisateur_reconciliation.docx"


def add_title(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(31, 78, 121)


def add_heading(document: Document, text: str, level: int = 1) -> None:
    document.add_heading(text, level=level)


def add_body(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(text)
    paragraph.paragraph_format.space_after = Pt(6)


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def add_capture(document: Document, filename: str, caption: str) -> None:
    image_path = SCREENSHOTS / filename
    if image_path.exists():
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run()
        run.add_picture(str(image_path), width=Cm(16.5))
        cap = document.add_paragraph(caption)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.runs[0].italic = True
        cap.runs[0].font.size = Pt(9)
    else:
        add_body(document, f"Capture indisponible: {filename}")


def add_kv_table(document: Document, rows: list[tuple[str, str]]) -> None:
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Element"
    table.rows[0].cells[1].text = "Description"
    for key, value in rows:
        cells = table.add_row().cells
        cells[0].text = key
        cells[1].text = value


def build_document() -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)

    styles = document.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(10)

    add_title(document, "Manuel utilisateur - Plateforme de reconciliation")
    subtitle = document.add_paragraph("Reconciliation Banque / Operateurs, Reporting, Compensation et Comptabilisation")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(12)
    document.add_paragraph()

    add_heading(document, "1. Objectif du document")
    add_body(
        document,
        "Ce manuel explique les principales actions a realiser dans l'application: se connecter, creer un compte, importer les fichiers, lancer une reconciliation, analyser le reporting, verifier la compensation et suivre la comptabilisation AMPLITUDE."
    )
    add_body(
        document,
        "Les exemples visuels sont des captures de demonstration. Les chiffres visibles servent a expliquer la lecture des ecrans; les utilisateurs doivent toujours se referer aux donnees reelles chargees dans leur environnement."
    )

    add_heading(document, "2. Connexion")
    add_capture(document, "01-connexion.png", "Ecran de connexion a l'application.")
    add_bullets(document, [
        "Saisir l'adresse email dans le champ identifiant. Dans l'application, l'email sert de nom d'utilisateur.",
        "Saisir le mot de passe puis cliquer sur Connexion.",
        "En cas d'echec, verifier l'email, le mot de passe et l'etat du compte utilisateur."
    ])

    add_heading(document, "3. Creation d'un compte")
    add_capture(document, "02-inscription.png", "Formulaire de creation de compte utilisateur.")
    add_body(document, "Chaque utilisateur peut creer son propre compte depuis le formulaire d'inscription.")
    add_kv_table(document, [
        ("Champs obligatoires", "Nom, prenom, email, mot de passe et confirmation du mot de passe."),
        ("Nom d'utilisateur", "L'email renseigne devient automatiquement le username."),
        ("Champs optionnels", "Les autres informations personnelles restent optionnelles selon la fiche utilisateur."),
        ("Controle attendu", "Le mot de passe et sa confirmation doivent etre identiques.")
    ])

    add_heading(document, "4. Accueil et navigation")
    add_capture(document, "03-accueil.png", "Page d'accueil apres authentification.")
    add_body(
        document,
        "La page d'accueil donne acces aux modules principaux: tableau de bord, reconciliation, compensation, comptabilisation et historique. L'utilisateur choisit le module selon l'activite a realiser."
    )

    add_heading(document, "5. Tableau de bord et reporting managerial")
    add_capture(document, "04-dashboard.png", "Tableau de bord avec indicateurs et filtres de reporting.")
    add_body(
        document,
        "Le tableau de bord presente une vision globale des operations: volumes, montants banque, montants operateur, anomalies, ecart net et distribution des resultats."
    )
    add_bullets(document, [
        "Utiliser le filtre Operateur pour basculer entre MOOV et ORANGE.",
        "Utiliser les filtres de date pour analyser une journee, une semaine ou un mois.",
        "Utiliser le filtre Type operation pour distinguer Banque vers Wallet et Wallet vers Banque.",
        "Le bloc Montant Anomalies doit etre lu avec les compteurs de transactions abouties cote operateur et cote Carthago.",
        "Les approvisionnements sont ecartes du rapprochement financier et doivent etre consultes dans leur fenetre dediee."
    ])

    add_heading(document, "6. Reconciliation")
    add_capture(document, "05-reconciliation.png", "Ecran de reconciliation et liste des resultats.")
    add_body(
        document,
        "Le module de reconciliation compare les donnees Banque/Carthago avec les donnees operateur. Il permet d'importer les fichiers, de lancer un run et de consulter les resultats detailles."
    )
    add_bullets(document, [
        "Importer les fichiers du meme perimetre metier: date, operateur et sens d'operation.",
        "Lancer la reconciliation apres verification des imports.",
        "Analyser les statuts: rapproche, absent cote banque, absent cote operateur, montant different, doublon ou approvisionnement.",
        "Utiliser les filtres pour isoler Banque vers Wallet, Wallet vers Banque, MOOV ou ORANGE.",
        "Les operations annulees ou rejetees ne doivent pas etre interpretees comme des transactions abouties."
    ])

    add_heading(document, "7. Compensation")
    add_capture(document, "06-compensation.png", "Vue compensation avec ecarts et decisions.")
    add_body(
        document,
        "La compensation sert a comparer les transactions abouties cote operateur avec les transactions abouties cote banque. Elle met en evidence l'ecart net, les volumes aboutis et les operations a justifier."
    )
    add_bullets(document, [
        "Selectionner l'operateur, la periode et le type d'operation.",
        "Verifier les compteurs aboutis operateur et aboutis Carthago avant d'analyser le montant.",
        "Consulter Operations a justifier pour les transactions abouties presentes d'un cote et absentes de l'autre.",
        "Consulter Ecarts a justifier pour les transactions abouties avec un ecart de montant.",
        "Les approvisionnements doivent etre sortis du calcul de compensation et controles separement."
    ])

    add_heading(document, "8. Comptabilisation AMPLITUDE")
    add_capture(document, "07-comptabilisation.png", "Controle de comptabilisation dans AMPLITUDE.")
    add_body(
        document,
        "Le module de comptabilisation permet de verifier si les operations Carthago attendues existent dans AMPLITUDE. Il fait ressortir les operations comptabilisees, non comptabilisees et les lignes AMPLITUDE sans correspondance Carthago."
    )
    add_bullets(document, [
        "Choisir la periode de controle et, si necessaire, l'operateur.",
        "Verifier le taux de comptabilisation en nombre et en montant.",
        "Analyser les operations non comptabilisees et les ecarts de montant.",
        "Exporter les controles si un traitement metier ou comptable est necessaire."
    ])

    add_heading(document, "9. Nettoyage et suppression en cascade")
    add_body(
        document,
        "Le nettoyage supprime des imports et peut impacter les transactions, les runs de reconciliation et les resultats deja calcules. Lorsque l'application signale une suppression en cascade, elle demande une confirmation explicite pour eviter une suppression involontaire."
    )
    add_bullets(document, [
        "Utiliser la preview de suppression avant confirmation.",
        "Lire le nombre d'imports, transactions, runs et resultats impactes.",
        "Confirmer la cascade uniquement si les donnees doivent reellement etre retirees du perimetre.",
        "Supprimer un import MOOV ou ORANGE ne supprime pas les donnees banque, sauf si une action de suppression banque est lancee separement."
    ])

    add_heading(document, "10. Bonnes pratiques")
    add_bullets(document, [
        "Toujours travailler avec le meme operateur, la meme date et le meme sens d'operation pour les imports compares.",
        "Relancer la reconciliation apres une evolution de regle metier, par exemple l'exclusion des approvisionnements.",
        "Controler les volumes avant les montants: un ecart de montant peut venir d'une operation absente.",
        "Conserver les exports de reporting et compensation pour validation metier.",
        "Isoler les approvisionnements avant toute interpretation des ecarts de compensation."
    ])

    DOCX_PATH.parent.mkdir(parents=True, exist_ok=True)
    document.save(DOCX_PATH)


if __name__ == "__main__":
    build_document()
    print(DOCX_PATH)
